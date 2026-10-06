from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
src=Path('C:/Users/welde/OneDrive/Imagens/Screenshots')
out=Path('C:/Users/welde/OneDrive/Área de Trabalho/Comet/30 exercicios.docx')
assert all((src/f'{i}{suffix}.png').is_file() for i in range(1,31) for suffix in ('','d'))
doc=Document()
s=doc.sections[0]
s.orientation=WD_ORIENT.LANDSCAPE
s.page_width=Inches(11.69); s.page_height=Inches(8.27)
s.top_margin=s.bottom_margin=Inches(.45)
s.left_margin=s.right_margin=Inches(.45)
for name in ('Normal','Title','Heading 1'):
 st=doc.styles[name]; st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
 st.paragraph_format.space_after=Pt(4)
doc.styles['Normal'].font.size=Pt(10)
doc.styles['Heading 1'].font.size=Pt(13)
doc.styles['Title'].font.size=Pt(18)
for i in range(1,31):
 if i%2==1 and i>1: doc.add_page_break()
 if i==1:
  doc.add_paragraph('30 exercícios de programação em blocos', 'Title')
 h=doc.add_paragraph(f'Exercício {i}', 'Heading 1')
 h.paragraph_format.keep_with_next=True
 table=doc.add_table(rows=2, cols=2)
 table.autofit=False; table.alignment=WD_TABLE_ALIGNMENT.CENTER
 widths=[2.55,8.14]
 for c,w in zip(table.columns,widths): c.width=Inches(w)
 for row in table.rows:
  trpr=row._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
  for cell,w in zip(row.cells,widths):
   cell.width=Inches(w); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcpr=cell._tc.get_or_add_tcPr()
   borders=OxmlElement('w:tcBorders')
   for edge in ('top','left','bottom','right'):
    e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
   tcpr.append(borders)
 for cell,label in zip(table.rows[0].cells,('Tela','Blocos')):
  p=cell.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run(label).bold=True
  shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'EFEFEF'); cell._tc.get_or_add_tcPr().append(shade)
 for col,suffix in enumerate(('d','')):
  path=src/f'{i}{suffix}.png'; iw,ih=Image.open(path).size
  maxw=widths[col]-.18; maxh=2.65
  scale=min(maxw/iw,maxh/ih)
  p=table.rows[1].cells[col].paragraphs[0]
  p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(4)
  shape=p.add_run().add_picture(str(path),width=Inches(iw*scale),height=Inches(ih*scale))
  shape._inline.docPr.set('descr',f'{"Tela" if col==0 else "Blocos"} do exercício {i}')
 gap=doc.add_paragraph(); gap.paragraph_format.space_after=Pt(0); gap.paragraph_format.space_before=Pt(0); gap.paragraph_format.line_spacing=Pt(3); gap.add_run().font.size=Pt(3)
doc.core_properties.title='30 exercícios de programação em blocos'
doc.save(out)
assert len(doc.inline_shapes)==60 and len(doc.tables)==30
print(out)
